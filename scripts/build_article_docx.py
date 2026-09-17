from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE = "A Multitask Transformer and Web System for Joint Sentiment Analysis and Review Authenticity Detection in E-Commerce: A Controlled Low-Resource Study"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a DOCX version of the English article draft.")
    parser.add_argument(
        "--input",
        type=Path,
        default=Path("docs/article_draft_en.md"),
        help="Markdown source file.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/article_draft_en.docx"),
        help="DOCX output path.",
    )
    return parser.parse_args()


def set_font(run, name: str, size: float, *, bold: bool = False, italic: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)


def set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{key}"))
        if element is None:
            element = OxmlElement(f"w:{key}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def apply_table_layout(table, column_count: int) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    usable_width = Inches(6.5)
    if column_count == 0:
        return
    widths = [usable_width / column_count] * column_count
    if column_count >= 5:
        widths[0] = Inches(1.6)
        widths[1] = Inches(1.1)
        remaining = usable_width - widths[0] - widths[1]
        shared = remaining / (column_count - 2)
        widths = widths[:2] + [shared] * (column_count - 2)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def ensure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_map = {
        "Heading 1": (16, "2E74B5", 16, 8),
        "Heading 2": (13, "2E74B5", 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for name, (size, color, before, after) in heading_map.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    if "Code Block" not in document.styles:
        style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.name = "Consolas"
        style.font.size = Pt(9.5)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.line_spacing = 1.05


INLINE_TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)")


def add_inline_runs(paragraph, text: str) -> None:
    for token in INLINE_TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, "Calibri", 11, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, "Consolas", 10)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, "Calibri", 11, italic=True)
        elif token.startswith("[") and "](" in token and token.endswith(")"):
            label, url = token[1:].split("](", 1)
            run = paragraph.add_run(f"{label} ({url[:-1]})")
            set_font(run, "Calibri", 11, color="1F4D78")
        else:
            run = paragraph.add_run(token.replace("``", '"'))
            set_font(run, "Calibri", 11)


def collect_paragraph(lines: list[str], start: int) -> tuple[str, int]:
    buffer: list[str] = []
    index = start
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            break
        if line.startswith(("# ", "## ", "### ", "|", "```", "- ", "1. ", "2. ", "3. ", "4. ", "5. ", "6. ", "7. ", "8. ", "9. ")):
            if buffer:
                break
        buffer.append(line.strip())
        index += 1
    return " ".join(buffer).strip(), index


def is_ordered_item(line: str) -> bool:
    return bool(re.match(r"^\d+\.\s+", line))


def build_manual_numbered_item(document: Document, number: int, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.35)
    paragraph.paragraph_format.first_line_indent = Inches(-0.35)
    paragraph.paragraph_format.space_after = Pt(4)
    marker = paragraph.add_run(f"{number}. ")
    set_font(marker, "Calibri", 11)
    add_inline_runs(paragraph, text)


def build_docx(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    ensure_styles(document)

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = document.add_paragraph(style="Normal")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing = 1.05
    add_inline_runs(title, TITLE)
    for run in title.runs:
        set_font(run, "Calibri", 18, bold=True, color="0B2545")

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped == f"# {TITLE}":
            index += 1
            continue

        if stripped.startswith("## "):
            heading = document.add_paragraph(style="Heading 1")
            add_inline_runs(heading, stripped[3:])
            index += 1
            continue

        if stripped.startswith("### "):
            heading = document.add_paragraph(style="Heading 2")
            add_inline_runs(heading, stripped[4:])
            index += 1
            continue

        if stripped.startswith("# "):
            heading = document.add_paragraph(style="Heading 1")
            add_inline_runs(heading, stripped[2:])
            index += 1
            continue

        if stripped.startswith("```"):
            fence_name = stripped[3:].strip()
            if fence_name:
                label = document.add_paragraph(style="Code Block")
                add_inline_runs(label, fence_name.upper())
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code = document.add_paragraph(style="Code Block")
                code.paragraph_format.space_after = Pt(0)
                run = code.add_run(lines[index])
                set_font(run, "Consolas", 9.5)
                index += 1
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|---"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                rows.append(row)
                index += 1
            header = rows[0]
            body_rows = rows[2:] if len(rows) > 2 else []
            table = document.add_table(rows=1, cols=len(header))
            apply_table_layout(table, len(header))
            for col, value in enumerate(header):
                cell = table.rows[0].cells[col]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline_runs(paragraph, value)
                shade_cell(cell, "F2F4F7")
                for run in paragraph.runs:
                    run.bold = True
            for row_values in body_rows:
                row = table.add_row().cells
                for col, value in enumerate(row_values):
                    row[col].text = ""
                    paragraph = row[col].paragraphs[0]
                    add_inline_runs(paragraph, value)
            document.add_paragraph()
            continue

        if stripped.startswith("- "):
            while index < len(lines) and lines[index].strip().startswith("- "):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(4)
                add_inline_runs(paragraph, lines[index].strip()[2:])
                index += 1
            continue

        if is_ordered_item(stripped):
            list_number = 1
            while index < len(lines) and is_ordered_item(lines[index].strip()):
                text = re.sub(r"^\d+\.\s+", "", lines[index].strip())
                build_manual_numbered_item(document, list_number, text)
                list_number += 1
                index += 1
            continue

        paragraph_text, index = collect_paragraph(lines, index)
        if paragraph_text:
            paragraph = document.add_paragraph(style="Normal")
            add_inline_runs(paragraph, paragraph_text)
            if paragraph_text.startswith("**Keywords:**"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    args = parse_args()
    build_docx(args.input, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
