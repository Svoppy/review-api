from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output/presentations/assignment5"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)

# Crop the live Stage View to the two completed runs, keeping the Jenkins UI intact.
stage_path = ASSETS / "jenkins-stage-view.png"
with Image.open(stage_path) as im:
    im.crop((0, 0, im.width, 445)).save(stage_path)
(ASSETS / "jenkins-stage-view-successes.png").unlink(missing_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

NAVY = "13283F"
NAVY2 = "203B55"
TEAL = "147D77"
MINT = "DDF3EE"
GOLD = "E7A64B"
CORAL = "D76B5C"
INK = "203041"
MUTED = "607180"
PAPER = "F7F8F6"
WHITE = "FFFFFF"
PALE = "EAF0F3"
LINE = "D6E0E5"
GREEN = "3F9467"
FONT = "Aptos"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_bg(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.02, italic=False, url=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    if url:
        run.hyperlink.address = url
    return box


def add_rich(slide, segments, x, y, w, h, size=18, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    for text, opts in segments:
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = opts.get("bold", False)
        r.font.italic = opts.get("italic", False)
        r.font.color.rgb = rgb(opts.get("color", color))
        if opts.get("url"):
            r.hyperlink.address = opts["url"]
    return box


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.2):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(line or fill)
    sh.line.width = Pt(0.8)
    if radius:
        try:
            sh.adjustments[0] = 0.08
        except Exception:
            pass
    return sh


def add_circle(slide, x, y, d, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(line or fill)
    return sh


def header(slide, title, section=None):
    add_text(slide, title, 0.62, 0.38, 12.05, 0.60, size=27, color=NAVY, bold=True)
    add_rect(slide, 0.63, 1.12, 0.78, 0.055, TEAL)


def footer(slide, number):
    add_line(slide, 0.62, 7.12, 12.72, 7.12, LINE, 0.6)
    add_text(slide, "REVIEWGUARD  /  DEVOPS & CI/CD", 0.66, 7.18, 4.5, 0.16, size=8, color=MUTED, bold=True)
    add_text(slide, f"{number:02d}", 12.12, 7.17, 0.5, 0.18, size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def new_slide(title=None, section=None, bg=PAPER):
    s = prs.slides.add_slide(BLANK)
    add_bg(s, bg)
    if title:
        header(s, title, section)
    return s


def bullet(slide, text, x, y, w, size=17, color=INK, accent=TEAL, h=0.52):
    add_circle(slide, x, y + 0.12, 0.10, accent)
    add_text(slide, text, x + 0.23, y, w - 0.23, h, size=size, color=color)


def stage_box(slide, x, y, w, title, body, accent=TEAL, number=None):
    add_rect(slide, x, y, w, 1.45, WHITE, LINE, radius=True)
    add_rect(slide, x, y, 0.06, 1.45, accent)
    title_y = y + 0.18
    if number:
        add_circle(slide, x + 0.20, y + 0.18, 0.30, accent)
        add_text(slide, str(number), x + 0.20, y + 0.20, 0.30, 0.16, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        title_x = x + 0.60
        title_w = w - 0.78
    else:
        title_x = x + 0.24
        title_w = w - 0.48
    add_text(slide, title, title_x, title_y, title_w, 0.34, size=15, color=NAVY, bold=True)
    add_text(slide, body, x + 0.24, y + 0.66, w - 0.46, 0.62, size=12, color=MUTED)


def add_picture_fit(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    nw, nh = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2), width=Inches(nw), height=Inches(nh))


def image_frame(slide, path, x, y, w, h):
    add_rect(slide, x, y, w, h, WHITE, LINE, radius=True)
    return add_picture_fit(slide, path, x + 0.06, y + 0.06, w - 0.12, h - 0.12)


def labeled_value(slide, label, value, x, y, w, accent=TEAL):
    add_rect(slide, x, y, 0.055, 0.66, accent)
    add_text(slide, value, x + 0.18, y - 0.01, w - 0.2, 0.32, size=22, color=NAVY, bold=True)
    add_text(slide, label, x + 0.18, y + 0.36, w - 0.2, 0.25, size=11, color=MUTED)


# 1 — cover
s = new_slide(bg=NAVY)
add_rect(s, 8.8, 0, 4.53, 7.5, NAVY2)
add_text(s, "ASSIGNMENT 5", 0.78, 0.72, 2.6, 0.3, size=12, color="91D4C8", bold=True)
add_text(s, "DevOps & CI/CD\nfor Research Projects", 0.76, 1.42, 7.45, 1.8, size=38, color=WHITE, bold=True)
add_text(s, "Automating a reproducible software workflow\nwith Docker and Jenkins", 0.80, 3.55, 6.5, 0.9, size=21, color="D8E4EB")
add_rect(s, 0.8, 4.72, 0.86, 0.055, TEAL)
add_text(s, "ReviewGuard  |  Python research API", 0.80, 5.05, 5.8, 0.4, size=15, color=WHITE, bold=True)
add_text(s, "Dias Kazikhanov", 0.80, 6.35, 3.5, 0.3, size=13, color="B9C8D4")
# simple pipeline motif
steps = [(9.3, 1.75, "Git"), (10.65, 2.55, "CI"), (9.3, 3.35, "Docker"), (10.65, 4.15, "Health")]
for i, (x, y, label) in enumerate(steps):
    add_circle(s, x, y, 0.62, TEAL if i < 3 else GOLD)
    add_text(s, str(i + 1), x, y + 0.17, 0.62, 0.22, size=14, color=WHITE if i < 3 else NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, label, x + 0.77, y + 0.16, 1.6, 0.26, size=14, color=WHITE, bold=True)
for a, b in [((9.61, 2.37), (10.85, 2.55)), ((10.85, 3.17), (9.61, 3.35)), ((9.61, 3.97), (10.85, 4.15))]:
    add_line(s, a[0], a[1], b[0], b[1], "A8BFCE", 1.6)
add_text(s, "A working CI demonstration,\nnot a claim of scientific validity", 9.3, 5.65, 3.5, 0.72, size=13, color="D8E4EB", italic=True)

# 2 — research reproducibility
s = new_slide("Research results depend on more than source code", "The research problem")
add_text(s, "A result is produced by a whole computational setup. If any part is missing, another researcher may not be able to rebuild the same workflow.", 0.66, 1.48, 11.9, 0.62, size=17, color=MUTED)
inputs = [
    (0.78, "CODE", "source revision\n& tests", TEAL),
    (3.76, "DATA", "dataset version\n& provenance", GOLD),
    (6.74, "CONFIG", "preprocessing\n& parameters", CORAL),
    (9.72, "RUNTIME", "packages, weights\n& hardware", NAVY2),
]
for x, title, body, color in inputs:
    add_rect(s, x, 2.60, 2.50, 1.40, WHITE, LINE, radius=True)
    add_rect(s, x, 2.60, 2.50, 0.12, color)
    add_text(s, title, x + 0.22, 2.92, 2.10, 0.28, size=14, color=color, bold=True)
    add_text(s, body, x + 0.22, 3.32, 2.08, 0.52, size=13, color=INK)
add_text(s, "DevOps makes the software steps visible and repeatable.", 0.82, 4.60, 11.5, 0.45, size=24, color=NAVY, bold=True)
add_text(s, "It records and checks how code is built and run. It does not decide whether a dataset, label scheme, or research claim is scientifically sound.", 0.84, 5.28, 11.2, 0.60, size=16, color=MUTED)
footer(s, 2)

# 3 — DevOps principles
s = new_slide("DevOps is a shared way of working", "Principles")
add_text(s, "The goal is a short, reliable feedback loop from a proposed change to a tested, runnable artifact.", 0.66, 1.48, 11.7, 0.44, size=17, color=MUTED)
items = [
    ("Version every change", "Keep code, pipeline definitions, and configuration in reviewable history.", TEAL),
    ("Automate routine checks", "Run the same tests and builds after every change.", GOLD),
    ("Describe environments", "Make runtime and deployment setup explicit instead of relying on a laptop state.", CORAL),
    ("Learn from failures", "Use fast, visible feedback to fix defects near the change that caused them.", NAVY2),
]
for i, (title, body, color) in enumerate(items):
    y = 2.18 + i * 1.05
    add_circle(s, 0.82, y + 0.02, 0.42, color)
    add_text(s, f"{i+1:02d}", 0.82, y + 0.12, 0.42, 0.15, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, title, 1.48, y, 2.65, 0.30, size=18, color=NAVY, bold=True)
    add_text(s, body, 4.42, y, 7.5, 0.54, size=15, color=INK)
add_rect(s, 0.82, 6.40, 11.6, 0.40, MINT)
add_text(s, "For research teams: fewer environment surprises, earlier regression detection, easier handoffs, and clearer provenance.", 1.02, 6.48, 11.2, 0.20, size=12, color=NAVY, bold=True)
footer(s, 3)

# 4 — CI/CD
s = new_slide("CI/CD turns each change into a repeatable check", "Workflow")
add_text(s, "Continuous integration (CI) and delivery (CD) share a pipeline, but they answer different questions.", 0.66, 1.48, 11.8, 0.45, size=17, color=MUTED)
stages = [
    (0.72, "Commit", "Git records\na change", NAVY2),
    (3.20, "Integrate", "CI runs tests\nand a build", TEAL),
    (5.68, "Package", "Docker creates\na runnable image", GOLD),
    (8.16, "Verify", "Smoke test the\ncontainer", CORAL),
    (10.64, "Deliver", "Promote to a test\nor release target", NAVY2),
]
for i, (x, title, body, color) in enumerate(stages):
    add_rect(s, x, 2.55, 1.90, 1.50, WHITE, LINE, radius=True)
    add_circle(s, x + 0.18, 2.76, 0.34, color)
    add_text(s, str(i + 1), x + 0.18, 2.85, 0.34, 0.15, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, title, x + 0.18, 3.20, 1.54, 0.25, size=15, color=NAVY, bold=True)
    add_text(s, body, x + 0.18, 3.55, 1.55, 0.45, size=11, color=MUTED)
    if i < len(stages) - 1:
        add_text(s, "›", x + 1.98, 3.09, 0.38, 0.42, size=29, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
add_text(s, "CI", 0.82, 4.55, 0.6, 0.32, size=15, color=TEAL, bold=True)
add_text(s, "Integrate small changes and verify them automatically.", 1.55, 4.55, 4.3, 0.34, size=15, color=INK)
add_text(s, "CD", 6.65, 4.55, 0.6, 0.32, size=15, color=GOLD, bold=True)
add_text(s, "Keep a tested version ready to release; deployment may remain gated.", 7.38, 4.55, 4.7, 0.52, size=15, color=INK)
add_rect(s, 0.82, 5.63, 11.5, 0.72, NAVY)
add_text(s, "In research software, a green pipeline proves the specified checks passed. It does not validate the scientific conclusion.", 1.08, 5.83, 10.9, 0.32, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
footer(s, 4)

# 5 — project architecture
s = new_slide("ReviewGuard pipeline: from Git to a health-checked container", "Implementation")
add_text(s, "The Jenkinsfile is versioned with the Python API and runs three operational stages after checkout.", 0.66, 1.48, 11.8, 0.45, size=17, color=MUTED)
architecture = [
    (0.78, "Git repository", "Python API\n+ Jenkinsfile", NAVY2),
    (3.25, "Jenkins", "Pipeline as Code\nlocal self-hosted", TEAL),
    (5.72, "Python agent", "Python 3.12\npytest", GOLD),
    (8.19, "Docker image", "reviewguard-ci:N\nsmall CI variant", CORAL),
    (10.66, "Smoke check", "temporary API\n+ /health", NAVY2),
]
for i, (x, title, body, color) in enumerate(architecture):
    add_rect(s, x, 2.42, 1.88, 1.32, WHITE, LINE, radius=True)
    add_rect(s, x, 2.42, 1.88, 0.09, color)
    add_text(s, title, x + 0.14, 2.70, 1.62, 0.30, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, body, x + 0.14, 3.10, 1.62, 0.46, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    if i < len(architecture) - 1:
        add_text(s, "→", x + 1.92, 2.89, 0.50, 0.30, size=22, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
add_line(s, 3.3, 4.35, 10.95, 4.35, LINE, 1.1)
add_text(s, "GitHub Actions also runs the full test suite on pushes and pull requests; Jenkins adds a local, Docker-aware demonstration path.", 1.02, 4.62, 11.3, 0.58, size=16, color=INK, align=PP_ALIGN.CENTER)
add_rect(s, 1.15, 5.65, 10.95, 0.70, PALE)
add_text(s, "Data and model checkpoints stay outside the image. Mount the trained checkpoint at /models/latest when inference is required.", 1.42, 5.86, 10.4, 0.30, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
footer(s, 5)

# 6 — Docker
s = new_slide("Docker packages the API runtime, not the research evidence", "Docker")
add_text(s, "The image makes the application easier to start consistently across a laptop, CI runner, and test host.", 0.66, 1.48, 11.8, 0.46, size=17, color=MUTED)
# left architecture stack
stack = [(0.92, 2.25, 4.8, 0.88, NAVY, "Python 3.12 slim base", "small, declared runtime"),
         (0.92, 3.24, 4.8, 0.88, TEAL, "ReviewGuard package + Uvicorn", "API listens on port 8000"),
         (0.92, 4.23, 4.8, 0.88, GOLD, "Non-root runtime user", "UID 10001  ·  explicit health check")]
for x, y, w, h, color, title, detail in stack:
    add_rect(s, x, y, w, h, WHITE, LINE, radius=True)
    add_rect(s, x, y, 0.08, h, color)
    add_text(s, title, x + 0.28, y + 0.16, w - 0.55, 0.28, size=16, color=NAVY, bold=True)
    add_text(s, detail, x + 0.28, y + 0.51, w - 0.55, 0.22, size=12, color=MUTED)
add_text(s, "→", 3.03, 5.25, 0.50, 0.35, size=24, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
add_rect(s, 1.45, 5.69, 3.75, 0.58, MINT)
add_text(s, "Container starts  →  /health responds", 1.58, 5.86, 3.48, 0.25, size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
# right explanatory notes
add_text(s, "KEEP OUT OF THE IMAGE", 6.65, 2.30, 5.5, 0.26, size=11, color=CORAL, bold=True)
bullet(s, "Raw and processed datasets", 6.65, 2.74, 5.3, size=15, h=0.35)
bullet(s, "Trained model weights", 6.65, 3.22, 5.3, size=15, h=0.35)
bullet(s, "Local reports and development files", 6.65, 3.70, 5.3, size=15, h=0.35)
add_text(s, "CI uses INSTALL_INFERENCE=0 for a smaller startup image. The default image installs inference dependencies; it still needs a checkpoint mounted separately.", 6.66, 4.52, 5.6, 1.05, size=14, color=INK)
add_text(s, "A container controls application packaging. Dataset versions, model provenance, seeds, and hardware still need their own records.", 6.66, 5.82, 5.55, 0.56, size=13, color=MUTED, italic=True)
footer(s, 6)

# 7 — Stage View screenshot
s = new_slide("Jenkins shows each stage and its result", "Pipeline evidence")
img = ASSETS / "jenkins-stage-view.png"
image_frame(s, img, 0.62, 1.48, 12.06, 3.82)
add_text(s, "BUILD #10", 0.85, 5.58, 1.3, 0.25, size=11, color=TEAL, bold=True)
add_text(s, "Checkout → unit tests → Docker image → container health check → post actions", 2.05, 5.55, 9.95, 0.31, size=15, color=NAVY, bold=True)
add_text(s, "Builds #10 and #11 are green. The earlier red rows show setup iterations visible in the same local Jenkins history.", 0.84, 6.12, 11.5, 0.50, size=13, color=MUTED)
footer(s, 7)

# 8 — local SCM run context
s = new_slide("Build #10 started from the local SCM snapshot", "Pipeline evidence")
image_frame(s, ASSETS / "jenkins-build-10-console.png", 0.62, 1.48, 8.30, 5.25)
add_text(s, "SOURCE REVISION", 9.28, 1.82, 3.2, 0.25, size=11, color=TEAL, bold=True)
add_text(s, "e9e91b1", 9.28, 2.20, 2.8, 0.50, size=28, color=NAVY, bold=True)
add_text(s, "The demonstration ran against a temporary local Git snapshot while Assignment 5 files were being prepared.", 9.28, 2.95, 3.25, 1.05, size=15, color=INK)
add_line(s, 9.28, 4.30, 12.45, 4.30, LINE, 1)
add_text(s, "The completed deliverables are now published on the repository's devops branch. This screenshot records the local demonstration run, not a post-push webhook run.", 9.28, 4.58, 3.25, 1.43, size=13, color=MUTED)
footer(s, 8)

# 9 — test result screenshot
s = new_slide("Automated tests give fast, repeatable feedback", "Pipeline evidence")
image_frame(s, ASSETS / "jenkins-build-10-tests.png", 0.62, 1.48, 8.35, 5.28)
add_text(s, "152", 9.42, 1.92, 2.2, 0.76, size=42, color=TEAL, bold=True)
add_text(s, "tests passed", 9.46, 2.78, 2.7, 0.35, size=19, color=NAVY, bold=True)
add_text(s, "3.64 s", 9.42, 3.48, 2.2, 0.50, size=26, color=NAVY, bold=True)
add_text(s, "test stage duration", 9.46, 4.02, 2.75, 0.30, size=12, color=MUTED)
add_rect(s, 9.42, 4.76, 2.76, 0.74, MINT)
add_text(s, "The Docker image build begins after pytest passes.", 9.60, 4.94, 2.40, 0.34, size=12, color=NAVY, bold=True)
add_text(s, "Evidence: Jenkins console output for build #10.", 9.42, 5.90, 2.9, 0.45, size=11, color=MUTED)
footer(s, 9)

# 10 — final container result
s = new_slide("The container starts and answers its health check", "Pipeline evidence")
image_frame(s, ASSETS / "jenkins-build-10-result.png", 0.62, 1.48, 8.30, 5.25)
add_text(s, "IMAGE", 9.28, 1.80, 2.9, 0.25, size=11, color=TEAL, bold=True)
add_text(s, "reviewguard-ci:10", 9.28, 2.16, 3.2, 0.44, size=19, color=NAVY, bold=True)
add_text(s, "HEALTH RESPONSE", 9.28, 2.95, 2.8, 0.25, size=11, color=TEAL, bold=True)
add_text(s, 'status: "ok"\nmodel_ready: false', 9.28, 3.30, 3.0, 0.80, size=17, color=NAVY, bold=True)
add_text(s, "The initial request raced container startup; a retry returned the response and the pipeline finished SUCCESS.", 9.28, 4.42, 3.2, 0.94, size=14, color=INK)
add_rect(s, 9.28, 5.70, 3.1, 0.70, PALE)
add_text(s, "Liveness verified; no trained checkpoint was loaded.", 9.48, 5.86, 2.70, 0.36, size=12, color=NAVY, bold=True)
footer(s, 10)

# 11 — evidence boundaries
s = new_slide("Operational checks are not scientific validation", "Interpretation")
add_text(s, "The pipeline provides evidence about software behavior at one recorded revision.", 0.66, 1.48, 11.8, 0.45, size=17, color=MUTED)
add_text(s, "WHAT THE RUN SUPPORTS", 0.88, 2.18, 5.4, 0.3, size=12, color=TEAL, bold=True)
add_text(s, "WHAT IT DOES NOT ESTABLISH", 6.92, 2.18, 5.4, 0.3, size=12, color=CORAL, bold=True)
add_line(s, 6.55, 2.10, 6.55, 6.40, LINE, 1.1)
for i, text in enumerate([
    "The specified test suite passed on the recorded source snapshot.",
    "The Docker image built from the declared project files.",
    "The API process started and /health returned a response.",
]):
    bullet(s, text, 0.90, 2.78 + i * 0.94, 5.2, size=16, h=0.60, accent=TEAL)
for i, text in enumerate([
    "Model accuracy or generalization to new review sources.",
    "Dataset representativeness or label validity.",
    "Bitwise reproducibility across hardware and runtimes.",
]):
    bullet(s, text, 6.94, 2.78 + i * 0.94, 5.2, size=16, h=0.60, accent=CORAL)
add_rect(s, 0.92, 5.95, 11.25, 0.57, NAVY)
add_text(s, "Research quality still depends on provenance for datasets, checkpoints, configuration, seeds, and source revisions.", 1.18, 6.12, 10.75, 0.23, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
footer(s, 11)

# 12 — implementation lessons
s = new_slide("Implementation exposed integration issues early", "Lessons")
add_text(s, "A useful pipeline does more than automate the happy path. It reveals mismatches between packaging, the runner, and the application.", 0.66, 1.48, 11.8, 0.55, size=16, color=MUTED)
# table headings and rows
add_text(s, "ISSUE", 0.90, 2.24, 3.1, 0.25, size=11, color=TEAL, bold=True)
add_text(s, "FIX / DESIGN CHOICE", 4.38, 2.24, 7.7, 0.25, size=11, color=TEAL, bold=True)
rows = [
    ("No active Docker daemon", "Start a local Docker runtime before building images."),
    ("Jenkins agent lacked Docker CLI / Git", "Add required tools to the Jenkins image and Python test agent."),
    ("API HTML asset missing from package", "Declare static HTML as package data; rerun the health check."),
    ("CPU-only CI runner", "Install CPU PyTorch wheels to avoid unused CUDA downloads."),
    ("Container startup race", "Retry /health briefly and always remove the temporary smoke container."),
]
for i, (issue, fix) in enumerate(rows):
    y = 2.68 + i * 0.70
    if i % 2 == 0:
        add_rect(s, 0.82, y - 0.08, 11.65, 0.62, WHITE, WHITE)
    add_text(s, issue, 0.95, y, 3.15, 0.35, size=13, color=NAVY, bold=True)
    add_text(s, fix, 4.40, y, 7.75, 0.40, size=13, color=INK)
add_text(s, "The health check now tests a real built container rather than only importing the API package.", 0.95, 6.35, 11.2, 0.32, size=14, color=TEAL, bold=True)
footer(s, 12)

# 13 — Kubernetes and conclusion
s = new_slide("Kubernetes is a deployment option, not a requirement", "Next steps")
add_text(s, "The assignment's working scope is Docker + Jenkins CI. Registry publishing and cluster deployment remain optional because no credentials or target cluster are configured.", 0.66, 1.48, 11.8, 0.66, size=16, color=MUTED)
# two-column plan
add_text(s, "A PRACTICAL NEXT PATH", 0.90, 2.52, 5.0, 0.3, size=12, color=TEAL, bold=True)
plan = [
    "Push a versioned image to a registry with scoped credentials.",
    "Add a webhook or multibranch job for the published devops branch.",
    "Deploy to a test environment and keep promotion gated.",
    "Use Kubernetes when shared scheduling or multiple services justify it.",
]
for i, item in enumerate(plan):
    bullet(s, item, 0.92, 3.05 + i * 0.67, 5.25, size=14, h=0.50, accent=TEAL)
add_line(s, 6.55, 2.48, 6.55, 5.85, LINE, 1.1)
add_text(s, "TAKEAWAY", 7.02, 2.52, 4.8, 0.3, size=12, color=GOLD, bold=True)
add_text(s, "DevOps makes research software easier to test, package, and rerun.", 7.02, 3.02, 5.0, 0.98, size=22, color=NAVY, bold=True)
add_text(s, "Use automation to strengthen traceability and reliability, while keeping scientific claims tied to data and empirical evidence.", 7.02, 4.32, 4.9, 0.95, size=15, color=INK)
add_rect(s, 0.92, 6.15, 11.2, 0.48, MINT)
add_rich(s, [("Repository: ", {"bold": True, "color": NAVY}), ("github.com/Svoppy/review-api/tree/devops", {"bold": True, "color": TEAL, "url": "https://github.com/Svoppy/review-api/tree/devops"})], 1.14, 6.27, 10.8, 0.23, size=13)
footer(s, 13)

# 14 — references
s = new_slide("References", "Sources")
refs = [
    ("Wilson et al. (2014). Best Practices for Scientific Computing. PLoS Biology 12(1): e1001745.", "https://doi.org/10.1371/journal.pbio.1001745"),
    ("Osborne et al. (2014). Ten simple rules for effective computational research. PLoS Comput. Biol. 10(1): e1003506.", "https://doi.org/10.1371/journal.pcbi.1003506"),
    ("DORA. Continuous Integration capability.", "https://dora.dev/capabilities/continuous-integration/"),
    ("DORA. Continuous Delivery capability.", "https://dora.dev/capabilities/continuous-delivery/"),
    ("Docker. Python language-specific guide.", "https://docs.docker.com/guides/python/"),
    ("Jenkins. Pipeline and Pipeline as Code documentation.", "https://www.jenkins.io/doc/book/pipeline/"),
    ("Kubernetes. Concepts documentation.", "https://kubernetes.io/docs/concepts/"),
]
for i, (citation, url) in enumerate(refs):
    y = 1.48 + i * 0.73
    add_circle(s, 0.88, y + 0.04, 0.26, TEAL if i < 2 else GOLD)
    add_text(s, str(i + 1), 0.88, y + 0.10, 0.26, 0.12, size=8, color=WHITE if i < 2 else NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, citation, 1.36, y, 10.9, 0.32, size=12, color=NAVY, bold=(i < 2))
    add_text(s, url, 1.38, y + 0.37, 10.7, 0.22, size=9, color=TEAL, url=url)
footer(s, 14)

pptx_path = OUT / "assignment5_devops_ci_cd.pptx"
prs.save(pptx_path)
print(f"Wrote {pptx_path}")
